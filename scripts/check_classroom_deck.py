#!/usr/bin/env python3
"""Classroom-oriented QA for editable PPTX decks.

Complements ``scripts/check_deck.py`` with checks that matter specifically on a
classroom projector: title orphaning, short-label wrapping, tiny text,
replacement/mojibake glyphs, slide overflow, and click-animation integrity.

The checker is deliberately conservative. Geometry heuristics WARN rather than
FAIL unless the failure is highly likely; final acceptance still requires a
render-to-PNG visual review.

Usage:
    python3 scripts/check_classroom_deck.py deck.pptx
    python3 scripts/check_classroom_deck.py deck.pptx --json qa.json
"""
from __future__ import annotations

import argparse
import json
import math
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation

EMU_PER_INCH = 914400
EMU_PER_PT = 12700
EXPECTED_W = 12192000
EXPECTED_H = 6858000
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"p": PML_NS}

TITLE_MIN_PT = 38.0
ABSOLUTE_MIN_PT = 20.0
BODY_TARGET_PT = 24.0
ENGLISH_TARGET_PT = 28.0
FOOTER_Y_IN = 6.80
MAX_RECOMMENDED_CLICKS = 6
MAX_ALLOWED_CLICKS = 8

SUSPICIOUS_GLYPHS = (
    ("Unicode replacement character", "\ufffd"),
    ("white square placeholder", "\u25a1"),
    ("double question mark", "??"),
)


def visual_units(text: str) -> float:
    """Approximate width in full-width-glyph units."""
    out = 0.0
    for ch in text:
        if ch in "\n\v\r":
            continue
        if ch.isspace():
            out += 0.35
        elif ord(ch) < 0x3000:
            out += 0.55
        else:
            out += 1.0
    return out


def visual_lines(text: str) -> list[str]:
    # PowerPoint uses vertical tab for manual soft line breaks.
    return [x for x in re.split(r"[\n\v\r]+", text) if x != ""] or [""]


def shape_text(shape) -> str:
    return (shape.text_frame.text or "") if getattr(shape, "has_text_frame", False) else ""


def run_font_sizes(shape) -> list[float]:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return sizes
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip() and r.font.size:
                sizes.append(float(r.font.size.pt))
    return sizes


def max_font_pt(shape) -> float | None:
    sizes = run_font_sizes(shape)
    return max(sizes) if sizes else None


def min_font_pt(shape) -> float | None:
    sizes = run_font_sizes(shape)
    return min(sizes) if sizes else None


def is_footer(shape, slide_h: int) -> bool:
    return float(shape.top) / EMU_PER_INCH >= FOOTER_Y_IN or shape.top >= int(slide_h * 0.88)


def is_small_kicker(shape, text: str) -> bool:
    top = float(shape.top) / EMU_PER_INCH
    size = max_font_pt(shape) or 0
    # e.g. PHRASAL VERBS / VQ II • 動詞② / Q1 / TAKEAWAY
    return top < 1.0 and size <= 22 and visual_units(text) <= 18


def choose_title_shape(slide):
    """Pick the classroom headline, not the tiny kicker or the hero sentence."""
    preferred = []
    fallback = []
    for shape in slide.shapes:
        text = shape_text(shape).strip()
        if not text:
            continue
        top = float(shape.top) / EMU_PER_INCH
        width = float(shape.width) / EMU_PER_INCH
        size = max_font_pt(shape) or 0.0
        if 0.45 <= top <= 1.45 and width >= 6.0 and size >= 30:
            preferred.append((top, -size, shape))
        elif top <= 1.55 and width >= 6.0 and size >= 30:
            fallback.append((top, -size, shape))
    pool = preferred or fallback
    return min(pool, default=(0, 0, None))[2]


def usable_width_pt(shape) -> float:
    tf = shape.text_frame
    width = float(shape.width) / EMU_PER_PT
    ml = float(tf.margin_left or 0) / EMU_PER_PT
    mr = float(tf.margin_right or 0) / EMU_PER_PT
    return max(1.0, width - ml - mr)


def predicted_line_count(shape, width_safety: float = 0.94) -> int:
    text = shape_text(shape).strip()
    if not text:
        return 0
    font = max_font_pt(shape) or 24.0
    cap = max(1.0, usable_width_pt(shape) * width_safety / font)
    total = 0
    for line in visual_lines(text):
        total += max(1, math.ceil(visual_units(line) / cap))
    return total


def predicted_tail_units(shape, width_safety: float = 0.94) -> float | None:
    text = shape_text(shape).strip()
    if not text or len(visual_lines(text)) > 1:
        return None
    font = max_font_pt(shape) or 24.0
    cap = max(1.0, usable_width_pt(shape) * width_safety / font)
    total = visual_units(text)
    lines = math.ceil(total / cap)
    if lines < 2:
        return None
    return total - cap * (lines - 1)


def japanese_ratio(text: str) -> float:
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return 0.0
    jp = sum(1 for c in chars if ord(c) >= 0x3000)
    return jp / len(chars)


def latin_ratio(text: str) -> float:
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return 0.0
    latin = sum(1 for c in chars if ord(c) < 128 and (c.isalpha() or c in "'.,!?;:-/()[]"))
    return latin / len(chars)


def inspect_animation_xml(pptx_path: Path, slide_no: int, fails: list[str], warns: list[str], result: dict):
    name = f"ppt/slides/slide{slide_no}.xml"
    with zipfile.ZipFile(pptx_path) as zf:
        if name not in zf.namelist():
            fails.append(f"p{slide_no}: slide XML is missing")
            return
        root = ET.fromstring(zf.read(name))

    shape_ids = {el.get("id") for el in root.findall(".//p:cNvPr", NS) if el.get("id")}
    timing = root.find("p:timing", NS)
    clicks = root.findall(".//p:cTn[@nodeType='clickEffect']", NS)
    result["click_effects"] = len(clicks)
    result["timing"] = timing is not None

    transition = root.find("p:transition", NS)
    if transition is not None and transition.get("advTm") is not None:
        fails.append(f"p{slide_no}: automatic slide advance is enabled (advTm); classroom decks must be presenter-controlled")

    if timing is None:
        return

    ids = [el.get("id") for el in timing.findall(".//p:cTn", NS) if el.get("id")]
    seen = set()
    dup = set()
    for x in ids:
        if x in seen:
            dup.add(x)
        seen.add(x)
    if dup:
        fails.append(f"p{slide_no}: duplicate animation time-node IDs: {sorted(dup)[:6]}")

    targets = [el.get("spid") for el in timing.findall(".//p:spTgt", NS) if el.get("spid")]
    missing = sorted({x for x in targets if x not in shape_ids})
    if missing:
        fails.append(f"p{slide_no}: animation targets refer to missing shapes: {missing}")

    if len(clicks) > MAX_ALLOWED_CLICKS:
        fails.append(f"p{slide_no}: {len(clicks)} click effects (> {MAX_ALLOWED_CLICKS}); split the slide")
    elif len(clicks) > MAX_RECOMMENDED_CLICKS:
        warns.append(f"p{slide_no}: {len(clicks)} click effects; recommended classroom range is 2–{MAX_RECOMMENDED_CLICKS}")


def run_checks(pptx_path: Path) -> dict:
    prs = Presentation(str(pptx_path))
    fails: list[str] = []
    warns: list[str] = []
    slide_summary = []

    if prs.slide_width != EXPECTED_W or prs.slide_height != EXPECTED_H:
        fails.append(f"canvas: expected 16:9 ({EXPECTED_W}x{EXPECTED_H} EMU), got {prs.slide_width}x{prs.slide_height}")

    for idx, slide in enumerate(prs.slides, start=1):
        sdata = {"slide": idx, "title": "", "click_effects": 0, "timing": False}
        title_shape = choose_title_shape(slide)
        if title_shape is not None:
            title = shape_text(title_shape).strip()
            sdata["title"] = title
            size = max_font_pt(title_shape)
            if idx != 1 and size is not None and size < TITLE_MIN_PT:
                fails.append(f"p{idx}: title is {size:.1f}pt (< {TITLE_MIN_PT:.0f}pt): {title[:55]!r}")

            explicit = visual_lines(title)
            if len(explicit) >= 2 and 0 < visual_units(explicit[-1].strip()) <= 4:
                fails.append(f"p{idx}: title ends with a lonely short line {explicit[-1].strip()!r}; rebreak at a semantic boundary")
            else:
                tail = predicted_tail_units(title_shape)
                # Automatic-wrap math is most trustworthy for Japanese-heavy titles;
                # Latin text width varies too much by font to hard-fail from geometry alone.
                if japanese_ratio(title) >= 0.65 and tail is not None and 0 < tail <= 4:
                    fails.append(f"p{idx}: title is very likely to auto-wrap with only ~{tail:.1f} full-width chars on the final line: {title[:60]!r}")
                elif japanese_ratio(title) >= 0.65 and tail is not None and tail <= 6:
                    warns.append(f"p{idx}: title may have a visually weak final line (~{tail:.1f} full-width chars): {title[:60]!r}")

        for shape in slide.shapes:
            text = shape_text(shape).strip()
            if not text:
                continue

            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                fails.append(f"p{idx}: text shape is outside slide bounds: {text[:50]!r}")

            for label, token in SUSPICIOUS_GLYPHS:
                if token in text:
                    fails.append(f"p{idx}: suspicious glyph ({label}): {text[:65]!r}")

            if is_footer(shape, prs.slide_height):
                continue

            minpt = min_font_pt(shape)
            if minpt is not None and minpt < ABSOLUTE_MIN_PT:
                fails.append(f"p{idx}: {minpt:.1f}pt text is below the {ABSOLUTE_MIN_PT:.0f}pt absolute classroom floor: {text[:55]!r}")
            elif minpt is not None and minpt < BODY_TARGET_PT and not is_small_kicker(shape, text):
                warns.append(f"p{idx}: {minpt:.1f}pt text is below the {BODY_TARGET_PT:.0f}pt classroom body target: {text[:55]!r}")

            if minpt is not None and latin_ratio(text) >= 0.75 and visual_units(text) > 20 and minpt < ENGLISH_TARGET_PT:
                warns.append(f"p{idx}: English-heavy text is {minpt:.1f}pt; important English should normally be ≥{ENGLISH_TARGET_PT:.0f}pt: {text[:55]!r}")

            # The most reliable geometry failure: a short label that should fit on one line
            # but its width/font ratio predicts wrapping. This catches cases like "take care of".
            if idx != 1 and len(visual_lines(text)) == 1 and visual_units(text) <= 18 and len(text.split()) <= 5 and not is_small_kicker(shape, text):
                font = max_font_pt(shape) or 24.0
                cap = max(1.0, usable_width_pt(shape) * 0.94 / font)
                ratio = visual_units(text) / cap
                if ratio > 1.15:
                    fails.append(f"p{idx}: short label is likely to wrap; width/font ratio={ratio:.2f}. Widen the box rather than shrink the font: {text!r}")
                elif ratio > 1.02:
                    warns.append(f"p{idx}: short label is close to wrapping; width/font ratio={ratio:.2f}. Check the rendered PNG: {text!r}")

            # Vertical geometry is renderer-dependent, so only WARN on very tight boxes.
            maxpt = max_font_pt(shape)
            if maxpt:
                available = float(shape.height) / EMU_PER_PT
                explicit_lines = len(visual_lines(text))
                if explicit_lines >= 2 and available < explicit_lines * maxpt * 0.92:
                    warns.append(f"p{idx}: multi-line text box may be vertically tight ({available:.0f}pt high for {explicit_lines}×{maxpt:.0f}pt): {text[:55]!r}")

        inspect_animation_xml(pptx_path, idx, fails, warns, sdata)
        slide_summary.append(sdata)

    return {
        "file": str(pptx_path),
        "slides": len(prs.slides),
        "fail_count": len(fails),
        "warn_count": len(warns),
        "fails": fails,
        "warnings": warns,
        "slide_summary": slide_summary,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--json", type=Path)
    args = ap.parse_args()

    if not args.pptx.exists():
        print(f"FAIL: file not found: {args.pptx}")
        return 2
    if args.pptx.suffix.lower() != ".pptx":
        print("FAIL: classroom checker currently accepts PPTX only")
        return 2

    report = run_checks(args.pptx)
    for msg in report["fails"]:
        print(f"FAIL: {msg}")
    for msg in report["warnings"]:
        print(f"WARN: {msg}")
    print(f"CLASSROOM_QA: slides={report['slides']} FAIL={report['fail_count']} WARN={report['warn_count']}")

    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 1 if report["fail_count"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
