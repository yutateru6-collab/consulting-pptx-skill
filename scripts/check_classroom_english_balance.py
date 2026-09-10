#!/usr/bin/env python3
"""Classroom QA for English line balance in PPTX.

Detects highly likely awkward line breaks in English-heavy text blocks.
This checker is intentionally conservative: explicit/manual breaks may FAIL;
renderer-dependent auto-wrap predictions are WARN and still require PNG review.
"""
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

from pptx import Presentation

EMU_PER_PT = 12700
FOOTER_Y_RATIO = 0.88
MIN_ENGLISH_UNITS = 14.0
MIN_WORDS = 4

BAD_BREAK_END = {
    "a", "an", "the", "to", "of", "for", "with", "at", "in", "on", "by",
    "from", "and", "or", "but", "if", "as", "than", "that", "which", "who",
    "would", "could", "might", "should", "will", "can", "have", "has", "had",
    "is", "are", "was", "were", "be", "been", "being",
}


def visual_units(text: str) -> float:
    out = 0.0
    for ch in text:
        if ch in "\n\v\r":
            continue
        if ch.isspace():
            out += 0.35
        elif ord(ch) < 128:
            out += 0.55
        else:
            out += 1.0
    return out


def latin_ratio(text: str) -> float:
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return 0.0
    latin = sum(1 for c in chars if ord(c) < 128 and (c.isalpha() or c in "'.,!?;:-/()[]"))
    return latin / len(chars)


def words(text: str) -> list[str]:
    return re.findall(r"[A-Za-z]+(?:'[A-Za-z]+)?", text)


def split_lines(text: str) -> list[str]:
    return [x.strip() for x in re.split(r"[\n\v\r]+", text) if x.strip()]


def min_font_pt(shape) -> float | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    vals = []
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip() and r.font.size:
                vals.append(float(r.font.size.pt))
    return min(vals) if vals else None


def max_font_pt(shape) -> float | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    vals = []
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip() and r.font.size:
                vals.append(float(r.font.size.pt))
    return max(vals) if vals else None


def usable_width_pt(shape) -> float:
    tf = shape.text_frame
    width = float(shape.width) / EMU_PER_PT
    ml = float(tf.margin_left or 0) / EMU_PER_PT
    mr = float(tf.margin_right or 0) / EMU_PER_PT
    return max(1.0, width - ml - mr)


def is_footer(shape, slide_h: int) -> bool:
    return shape.top >= int(slide_h * FOOTER_Y_RATIO)


def looks_like_english_sentence(text: str) -> bool:
    ws = words(text)
    return (
        latin_ratio(text) >= 0.68
        and visual_units(text) >= MIN_ENGLISH_UNITS
        and len(ws) >= MIN_WORDS
    )


def explicit_balance_issue(text: str) -> tuple[str | None, str | None]:
    """Return (severity, reason) for manual/explicit line breaks."""
    lines = split_lines(text)
    if len(lines) < 2 or not looks_like_english_sentence(text):
        return None, None

    # Only judge lines that are substantially English; this avoids Japanese translations
    # in a mixed text box dominating the balance calculation.
    eng_lines = [ln for ln in lines if latin_ratio(ln) >= 0.62 and words(ln)]
    if len(eng_lines) < 2:
        return None, None

    widths = [visual_units(ln) for ln in eng_lines]
    counts = [len(words(ln)) for ln in eng_lines]
    max_w = max(widths)
    min_w = min(widths)
    min_i = widths.index(min_w)
    min_ratio = min_w / max_w if max_w else 1.0

    # Very short orphan line such as "If I" or "the truth."
    if counts[min_i] <= 2 and min_w <= 12 and min_ratio < 0.55:
        return "FAIL", f"English orphan line: {eng_lines[min_i]!r} is much shorter than adjacent line(s)"

    # Very strong visual imbalance. Keep threshold conservative to avoid penalizing
    # semantically natural breaks that are merely a little uneven.
    if min_ratio < 0.38 and min_w <= 14:
        return "FAIL", f"English line balance is too uneven (short/long ratio={min_ratio:.2f})"

    # Function-word break: e.g. line ends in 'to', 'the', 'would'.
    for i, ln in enumerate(eng_lines[:-1]):
        ws = words(ln)
        if ws and ws[-1].lower() in BAD_BREAK_END:
            return "FAIL", f"Awkward English break after function word {ws[-1]!r}: {ln!r}"

    # Borderline imbalance is a warning and must be confirmed on rendered PNG.
    if min_ratio < 0.55 and min_w <= 18:
        return "WARN", f"English line balance may be visually weak (short/long ratio={min_ratio:.2f})"

    return None, None


def predicted_wrap_warning(shape, text: str) -> str | None:
    """Predict a likely ragged automatic 2-line wrap from box width/font size.

    This is renderer-dependent, so it never returns FAIL.
    """
    if len(split_lines(text)) != 1 or not looks_like_english_sentence(text):
        return None
    font = max_font_pt(shape) or min_font_pt(shape)
    if not font or font <= 0:
        return None

    # Same full-width-unit approximation used by classroom layout QA.
    capacity = usable_width_pt(shape) * 0.92 / font
    total = visual_units(text)
    if capacity <= 0:
        return None
    predicted = math.ceil(total / capacity)
    if predicted != 2:
        return None

    tail = total - capacity
    if tail <= 0:
        return None
    tail_ratio = tail / capacity
    if tail_ratio < 0.28:
        return f"likely ragged auto-wrap: predicted second line only {tail_ratio:.0%} of first-line capacity"
    return None


def run_checks(path: Path) -> dict:
    prs = Presentation(str(path))
    fails: list[str] = []
    warns: list[str] = []
    issues: list[dict] = []

    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = (shape.text_frame.text or "").strip()
            if not text or is_footer(shape, prs.slide_height):
                continue

            sev, reason = explicit_balance_issue(text)
            if sev:
                rec = {
                    "slide": slide_no,
                    "shape": getattr(shape, "name", ""),
                    "severity": sev.lower(),
                    "reason": reason,
                    "text": text[:220],
                }
                issues.append(rec)
                msg = f"p{slide_no}: {reason}: {text[:90]!r}"
                if sev == "FAIL":
                    fails.append(msg)
                else:
                    warns.append(msg)

            pred = predicted_wrap_warning(shape, text)
            if pred:
                issues.append({
                    "slide": slide_no,
                    "shape": getattr(shape, "name", ""),
                    "severity": "warn",
                    "reason": pred,
                    "text": text[:220],
                })
                warns.append(f"p{slide_no}: {pred}: {text[:90]!r}")

    return {
        "file": str(path),
        "slides": len(prs.slides),
        "fail_count": len(fails),
        "warn_count": len(warns),
        "fails": fails,
        "warnings": warns,
        "issues": issues,
        "note": "Auto-wrap predictions are heuristic; final acceptance still requires rendered-PNG review.",
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--json", type=Path)
    args = ap.parse_args()

    if not args.pptx.exists() or args.pptx.suffix.lower() != ".pptx":
        print(f"FAIL: invalid PPTX path: {args.pptx}")
        return 2

    result = run_checks(args.pptx)
    for x in result["fails"]:
        print("FAIL:", x)
    for x in result["warnings"]:
        print("WARN:", x)
    print(
        f"ENGLISH_BALANCE: slides={result['slides']} "
        f"FAIL={result['fail_count']} WARN={result['warn_count']}"
    )

    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    return 1 if result["fail_count"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
