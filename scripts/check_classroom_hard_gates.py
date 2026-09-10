#!/usr/bin/env python3
"""Hard-gate checker for classroom PPTX.

This is intentionally stricter than check_classroom_deck.py. It enforces
non-negotiable classroom delivery rules: projector-sized text, substantive
speaker notes, click-reveal coverage, borderless text containers, and
flowchart-specific anti-card-grid guardrails.

Usage:
  python3 scripts/check_classroom_hard_gates.py deck.pptx --json hard-gates.json
  python3 scripts/check_classroom_hard_gates.py deck.pptx --profile flowchart
  python3 scripts/check_classroom_hard_gates.py deck.pptx --allow-static --notes-optional
"""
from __future__ import annotations

import argparse
import json
import posixpath
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from classroom_style_policy import violates_borderless_policy

EMU = 914400
EXPECTED_W, EXPECTED_H = 12192000, 6858000
TOL = 2000
COVER_TITLE_MIN = 54.0
TITLE_MIN = 38.0
LABEL_MIN = 22.0
BODY_MIN = 24.0
ENGLISH_MIN = 28.0
FOOTER_Y = 6.80
PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
DML = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"p": PML, "a": DML}
FLOW_MARKERS = ("flowchart", "flow chart", "フローチャート", "判断フロー", "decision tree")
Q_MARKERS = ("演習", "練習", "quiz", "check", "問題", "問い", "question")
BAD_GLYPHS = ("\ufffd", "\u25a1", "??")


def txt(sh):
    return (sh.text_frame.text or "") if getattr(sh, "has_text_frame", False) else ""


def sizes(sh):
    return [
        r.font.size.pt
        for p in sh.text_frame.paragraphs
        for r in p.runs
        if r.text.strip() and r.font.size
    ] if getattr(sh, "has_text_frame", False) else []


def minpt(sh):
    s = sizes(sh)
    return min(s) if s else None


def maxpt(sh):
    s = sizes(sh)
    return max(s) if s else None


def footer(sh, prs):
    return sh.top / EMU >= FOOTER_Y or sh.top >= int(prs.slide_height * .88)


def units(s):
    n = 0.0
    for c in s:
        if c in "\n\v\r":
            continue
        if c.isspace():
            n += .35
        elif ord(c) < 0x3000:
            n += .55
        else:
            n += 1
    return n


def latin_ratio(s):
    cs = [c for c in s if not c.isspace()]
    if not cs:
        return 0
    return sum(1 for c in cs if ord(c) < 128 and (c.isalpha() or c in "'.,!?;:-/()[]+")) / len(cs)


def short_label(sh, s):
    if sh.top / EMU < 1.0 and units(s) <= 18:
        return True
    return "\n" not in s and "\v" not in s and units(s) <= 18 and len(s.split()) <= 5 and len(s) <= 34


def choose_title(slide, cover=False):
    cs = []
    for sh in slide.shapes:
        s = txt(sh).strip()
        if not s:
            continue
        top = sh.top / EMU
        width = sh.width / EMU
        size = maxpt(sh) or 0
        if cover:
            if top <= 4.5 and width >= 3.0:
                cs.append((-size, top, sh))
        else:
            if top <= 1.6 and width >= 4.5 and size >= 20:
                cs.append((top, -size, sh))
    return min(cs, default=(0, 0, None))[2]


def notes_text(zf, n):
    rel = f"ppt/slides/_rels/slide{n}.xml.rels"
    if rel not in zf.namelist():
        return ""
    rr = ET.fromstring(zf.read(rel))
    target = None
    for r in rr:
        if r.get("Type", "").endswith("/notesSlide"):
            target = r.get("Target")
            break
    if not target:
        return ""
    part = posixpath.normpath(posixpath.join("ppt/slides", target))
    if part not in zf.namelist():
        return ""
    root = ET.fromstring(zf.read(part))
    s = " ".join(
        (e.text or "").strip()
        for e in root.findall(".//a:t", NS)
        if (e.text or "").strip()
    )
    s = re.sub(r"\b\d+\b", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def click_count(zf, n):
    part = f"ppt/slides/slide{n}.xml"
    if part not in zf.namelist():
        return 0, False, False
    root = ET.fromstring(zf.read(part))
    clicks = len(root.findall(".//p:cTn[@nodeType='clickEffect']", NS))
    timing = root.find("p:timing", NS) is not None
    tr = root.find("p:transition", NS)
    auto = tr is not None and tr.get("advTm") is not None
    return clicks, timing, auto


def profile_for(path, prs, requested):
    if requested != "auto":
        return requested
    probe = path.name.lower()
    if prs.slides:
        probe += " " + " ".join(txt(s) for s in prs.slides[0].shapes).lower()
    return "flowchart" if any(m in probe for m in FLOW_MARKERS) else "standard"


def question_slide(slide):
    s = " ".join(txt(sh) for sh in slide.shapes).lower()
    return any(m in s for m in Q_MARKERS)


def check(path, profile="auto", allow_static=False, notes_optional=False):
    prs = Presentation(str(path))
    fails = []
    warns = []
    slides = []

    if abs(prs.slide_width - EXPECTED_W) > TOL or abs(prs.slide_height - EXPECTED_H) > TOL:
        fails.append(f"canvas is not 16:9 within tolerance: {prs.slide_width}x{prs.slide_height}")

    prof = profile_for(path, prs, profile)
    total_clicks = 0
    note_count = 0
    click_slides = 0

    with zipfile.ZipFile(path) as zf:
        for i, slide in enumerate(prs.slides, 1):
            rec = {"slide": i, "title": "", "clicks": 0, "notes": False}
            title = choose_title(slide, cover=(i == 1))

            if title is None:
                fails.append(f"p{i}: no usable classroom title found")
            else:
                t = txt(title).strip()
                rec["title"] = t
                s = maxpt(title) or 0
                need = COVER_TITLE_MIN if i == 1 else TITLE_MIN
                if s < need:
                    fails.append(f"p{i}: title {s:.1f}pt < {need:.0f}pt: {t[:60]!r}")
                if violates_borderless_policy(title):
                    fails.append(
                        f"p{i}: title uses a decorative rectangular outline; "
                        f"Classroom Mode text containers must be borderless: {t[:60]!r}"
                    )

            note = notes_text(zf, i)
            has = len(note) >= 12
            rec["notes"] = has
            if has:
                note_count += 1
            elif i != 1 and not notes_optional:
                fails.append(f"p{i}: substantive speaker notes missing")

            for sh in slide.shapes:
                s = txt(sh).strip()
                if not s or footer(sh, prs) or sh is title:
                    continue

                if violates_borderless_policy(sh):
                    fails.append(
                        f"p{i}: text-bearing rectangle has a decorative outline; "
                        f"use no-line text containers. Semantic borders must be explicitly "
                        f"named FLOW_/NODE_/TABLE_/AXIS_/VENN_/DIAGRAM_/UI_: {s[:60]!r}"
                    )

                if any(g in s for g in BAD_GLYPHS):
                    fails.append(f"p{i}: suspicious replacement glyph in {s[:60]!r}")

                p = minpt(sh)
                if p is None:
                    continue
                if short_label(sh, s):
                    if p < LABEL_MIN:
                        fails.append(f"p{i}: label {p:.1f}pt < {LABEL_MIN:.0f}pt: {s[:60]!r}")
                elif p < BODY_MIN:
                    fails.append(f"p{i}: body {p:.1f}pt < {BODY_MIN:.0f}pt: {s[:60]!r}")
                if latin_ratio(s) >= .75 and units(s) > 20 and p < ENGLISH_MIN:
                    fails.append(f"p{i}: English text {p:.1f}pt < {ENGLISH_MIN:.0f}pt: {s[:60]!r}")

            c, timing, auto = click_count(zf, i)
            rec["clicks"] = c
            total_clicks += c
            if c:
                click_slides += 1
            if auto:
                fails.append(f"p{i}: automatic slide advance is enabled")
            if i != 1 and not allow_static:
                if c == 0:
                    fails.append(f"p{i}: no on-click reveal effects")
                if question_slide(slide) and c < 2:
                    fails.append(f"p{i}: exercise/quiz slide needs >=2 click effects")
            slides.append(rec)

    if len(prs.slides) > 1 and not notes_optional:
        content = len(prs.slides) - 1
        n = sum(1 for r in slides[1:] if r["notes"])
        if n < content:
            fails.append(f"notes coverage {n}/{content}; expected every content slide")

    if len(prs.slides) > 1 and not allow_static:
        content = len(prs.slides) - 1
        c = sum(1 for r in slides[1:] if r["clicks"] > 0)
        if total_clicks == 0:
            fails.append("deck-level animation gate: 0 click effects")
        if c < content:
            fails.append(f"click coverage {c}/{content}; expected every content slide")

    if prof == "flowchart":
        flow = []
        cards = []
        for i, sl in enumerate(prs.slides, 1):
            lines = sum(1 for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.LINE)
            auto = sum(1 for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
            textblocks = sum(1 for sh in sl.shapes if txt(sh).strip())
            if lines >= 3 and auto >= 4:
                flow.append(i)
            if auto >= 8 and lines <= 1 and textblocks >= 8:
                cards.append(i)
        if len(prs.slides) >= 8 and len(flow) < 2:
            fails.append("flowchart profile: fewer than two structurally flow-like slides")
        if len(cards) >= 3:
            fails.append(
                f"flowchart profile: repeated card-grid slides {cards[:10]}; "
                f"use a persistent decision spine/branch zoom"
            )

    return {
        "file": str(path),
        "slides": len(prs.slides),
        "profile": prof,
        "fail_count": len(fails),
        "warn_count": len(warns),
        "fails": fails,
        "warnings": warns,
        "slide_summary": slides,
        "summary": {
            "total_click_effects": total_clicks,
            "slides_with_clicks": click_slides,
            "notes_present": note_count,
        },
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--json", type=Path)
    ap.add_argument("--profile", choices=("auto", "standard", "flowchart"), default="auto")
    ap.add_argument("--allow-static", action="store_true")
    ap.add_argument("--notes-optional", action="store_true")
    a = ap.parse_args()

    if not a.pptx.exists() or a.pptx.suffix.lower() != ".pptx":
        print(f"FAIL: invalid PPTX path: {a.pptx}")
        return 2

    result = check(a.pptx, a.profile, a.allow_static, a.notes_optional)
    for item in result["fails"]:
        print("FAIL:", item)
    for item in result["warnings"]:
        print("WARN:", item)
    print(
        f"HARD_GATES: profile={result['profile']} slides={result['slides']} "
        f"FAIL={result['fail_count']} WARN={result['warn_count']} "
        f"clicks={result['summary']['total_click_effects']} "
        f"notes={result['summary']['notes_present']}"
    )

    if a.json:
        a.json.parent.mkdir(parents=True, exist_ok=True)
        a.json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    return 1 if result["fail_count"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
