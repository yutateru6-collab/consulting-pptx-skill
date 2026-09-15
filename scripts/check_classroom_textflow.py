#!/usr/bin/env python3
"""Detect likely text overflow and downstream collisions in classroom PPTX decks.

This checker targets a failure that ordinary shape-overlap checks miss:
a text box itself ends before the next box, but the text auto-wraps to an
extra line and visually overflows into the following text block.

Usage:
  python3 scripts/check_classroom_textflow.py deck.pptx
  python3 scripts/check_classroom_textflow.py deck.pptx --json textflow.json
"""
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

from pptx import Presentation

EMU_PER_INCH = 914400
EMU_PER_PT = 12700
WIDTH_SAFETY = 0.91
LINE_HEIGHT_FACTOR = 1.18
EXTRA_VERTICAL_PT = 4.0
SAFE_GAP_IN = 0.14
HARD_HEIGHT_RATIO = 0.91
WARN_HEIGHT_RATIO = 1.00


def text_of(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return (shape.text_frame.text or "").strip()


def font_sizes(shape) -> list[float]:
    out: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return out
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip() and r.font.size:
                out.append(float(r.font.size.pt))
    return out


def max_font_pt(shape) -> float:
    vals = font_sizes(shape)
    return max(vals) if vals else 24.0


def visual_units(text: str) -> float:
    total = 0.0
    for ch in text:
        if ch in "\n\v\r":
            continue
        if ch.isspace():
            total += 0.35
        elif ord(ch) < 0x3000:
            total += 0.55
        else:
            total += 1.0
    return total


def explicit_lines(text: str) -> list[str]:
    return [x for x in re.split(r"[\n\v\r]+", text) if x != ""] or [""]


def usable_width_pt(shape) -> float:
    tf = shape.text_frame
    width = float(shape.width) / EMU_PER_PT
    ml = float(tf.margin_left or 0) / EMU_PER_PT
    mr = float(tf.margin_right or 0) / EMU_PER_PT
    return max(1.0, width - ml - mr)


def predicted_line_count(shape) -> int:
    text = text_of(shape)
    if not text:
        return 0
    font = max_font_pt(shape)
    capacity = max(1.0, usable_width_pt(shape) * WIDTH_SAFETY / font)
    total = 0
    for line in explicit_lines(text):
        total += max(1, math.ceil(visual_units(line) / capacity))
    return total


def vertical_margins_pt(shape) -> float:
    tf = shape.text_frame
    mt = float(tf.margin_top or 0) / EMU_PER_PT
    mb = float(tf.margin_bottom or 0) / EMU_PER_PT
    return mt + mb


def required_height_pt(shape) -> float:
    lines = predicted_line_count(shape)
    if lines <= 0:
        return 0.0
    return lines * max_font_pt(shape) * LINE_HEIGHT_FACTOR + vertical_margins_pt(shape) + EXTRA_VERTICAL_PT


def box_height_pt(shape) -> float:
    return float(shape.height) / EMU_PER_PT


def left_in(shape) -> float:
    return float(shape.left) / EMU_PER_INCH


def top_in(shape) -> float:
    return float(shape.top) / EMU_PER_INCH


def width_in(shape) -> float:
    return float(shape.width) / EMU_PER_INCH


def height_in(shape) -> float:
    return float(shape.height) / EMU_PER_INCH


def horizontal_overlap_in(a, b) -> float:
    l = max(left_in(a), left_in(b))
    r = min(left_in(a) + width_in(a), left_in(b) + width_in(b))
    return max(0.0, r - l)


def is_footer(shape, slide_h_in: float) -> bool:
    return top_in(shape) >= slide_h_in * 0.88


def latin_ratio(text: str) -> float:
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return 0.0
    latin = sum(1 for c in chars if ord(c) < 128 and (c.isalpha() or c in "'.,!?;:-/()[]"))
    return latin / len(chars)


def run_checks(path: Path) -> dict:
    prs = Presentation(str(path))
    fails: list[str] = []
    warns: list[str] = []
    slide_results: list[dict] = []
    slide_h_in = float(prs.slide_height) / EMU_PER_INCH

    for slide_no, slide in enumerate(prs.slides, start=1):
        shapes = [s for s in slide.shapes if text_of(s) and not is_footer(s, slide_h_in)]
        details = []

        for s in shapes:
            text = text_of(s)
            lines = predicted_line_count(s)
            req = required_height_pt(s)
            avail = box_height_pt(s)
            ratio = avail / req if req > 0 else 99.0
            details.append({
                "text": text[:120],
                "predicted_lines": lines,
                "font_pt": round(max_font_pt(s), 1),
                "box_height_pt": round(avail, 1),
                "required_height_pt": round(req, 1),
                "height_ratio": round(ratio, 3),
            })

            if req > 0 and ratio < HARD_HEIGHT_RATIO:
                fails.append(
                    f"p{slide_no}: likely text overflow: box {avail:.0f}pt < estimated {req:.0f}pt "
                    f"for {lines} line(s) at {max_font_pt(s):.0f}pt: {text[:70]!r}"
                )
            elif req > 0 and ratio < WARN_HEIGHT_RATIO:
                warns.append(
                    f"p{slide_no}: text box has almost no vertical reserve ({avail:.0f}pt vs est. {req:.0f}pt): {text[:70]!r}"
                )

            if latin_ratio(text) >= 0.72 and max_font_pt(s) >= 28 and width_in(s) <= 4.25 and lines >= 3:
                warns.append(
                    f"p{slide_no}: large English in a narrow column needs {lines} lines; consider 3→2 columns or a shorter example: {text[:70]!r}"
                )

        # Detect the key failure mode: geometric boxes do not overlap, but estimated
        # rendered text from the upper shape reaches the following text block.
        for a in shapes:
            req_in = required_height_pt(a) / 72.0
            predicted_bottom = top_in(a) + req_in
            geometric_bottom = top_in(a) + height_in(a)
            if predicted_bottom <= geometric_bottom + 0.01:
                continue
            for b in shapes:
                if a is b or top_in(b) <= top_in(a):
                    continue
                if horizontal_overlap_in(a, b) < min(width_in(a), width_in(b)) * 0.25:
                    continue
                # Only flag when the authored rectangles themselves were separated.
                if geometric_bottom <= top_in(b) and predicted_bottom + SAFE_GAP_IN > top_in(b):
                    fails.append(
                        f"p{slide_no}: overflow from upper text is likely to collide with the next block; "
                        f"estimated bottom={predicted_bottom:.2f}in, next top={top_in(b):.2f}in, "
                        f"required gap={SAFE_GAP_IN:.2f}in. Upper={text_of(a)[:45]!r} / lower={text_of(b)[:45]!r}"
                    )

        slide_results.append({"slide": slide_no, "text_shapes": details})

    return {
        "file": str(path),
        "fail_count": len(fails),
        "warn_count": len(warns),
        "fails": fails,
        "warnings": warns,
        "slides": slide_results,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--json", type=Path)
    args = ap.parse_args()

    if not args.pptx.exists():
        print(f"FAIL: file not found: {args.pptx}")
        return 2
    if args.pptx.suffix.lower() != ".pptx":
        print("FAIL: expected a .pptx file")
        return 2

    result = run_checks(args.pptx)
    if args.json:
        args.json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    for item in result["fails"]:
        print("FAIL", item)
    for item in result["warnings"]:
        print("WARN", item)
    print(f"textflow: FAIL={result['fail_count']} WARN={result['warn_count']}")
    return 1 if result["fail_count"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
