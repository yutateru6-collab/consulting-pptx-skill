#!/usr/bin/env python3
"""Classroom visual-asset QA for PPTX decks.

Checks picture placement, text overlap, aspect-ratio distortion, and visual density.
This complements rendered PNG review; it does not replace it.

Usage:
    python3 scripts/check_classroom_visual_assets.py deck.pptx
    python3 scripts/check_classroom_visual_assets.py deck.pptx --json visual-assets-qa.json
"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_INCH = 914400
TEXT_GAP_FAIL_IN = 0.04
TEXT_GAP_WARN_IN = 0.12
SIGNIFICANT_IMAGE_AREA_RATIO = 0.025
MAX_SIGNIFICANT_IMAGES = 2
DISTORTION_FAIL_RATIO = 0.10
OVERLAY_ALLOWED_PREFIXES = ("BG_", "FULLBLEED_", "DIAGRAM_BG_")


def _rect(shape) -> tuple[float, float, float, float]:
    x1 = float(shape.left) / EMU_PER_INCH
    y1 = float(shape.top) / EMU_PER_INCH
    x2 = float(shape.left + shape.width) / EMU_PER_INCH
    y2 = float(shape.top + shape.height) / EMU_PER_INCH
    return x1, y1, x2, y2


def _area(rect: tuple[float, float, float, float]) -> float:
    x1, y1, x2, y2 = rect
    return max(0.0, x2 - x1) * max(0.0, y2 - y1)


def _intersection(a, b) -> float:
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    w = max(0.0, min(ax2, bx2) - max(ax1, bx1))
    h = max(0.0, min(ay2, by2) - max(ay1, by1))
    return w * h


def _projection_overlap(a1: float, a2: float, b1: float, b2: float) -> bool:
    return min(a2, b2) > max(a1, b1)


def _edge_gap(a, b) -> float | None:
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    if _intersection(a, b) > 0:
        return 0.0

    candidates: list[float] = []
    if _projection_overlap(ay1, ay2, by1, by2):
        if ax2 <= bx1:
            candidates.append(bx1 - ax2)
        elif bx2 <= ax1:
            candidates.append(ax1 - bx2)
    if _projection_overlap(ax1, ax2, bx1, bx2):
        if ay2 <= by1:
            candidates.append(by1 - ay2)
        elif by2 <= ay1:
            candidates.append(ay1 - by2)

    return min(candidates) if candidates else None


def _shape_name(shape) -> str:
    return str(getattr(shape, "name", "") or "")


def _text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return str(shape.text_frame.text or "").strip()


def _is_picture(shape) -> bool:
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE


def _overlay_allowed(shape) -> bool:
    name = _shape_name(shape).upper()
    return any(name.startswith(prefix) for prefix in OVERLAY_ALLOWED_PREFIXES)


def _crop_is_zero(shape) -> bool:
    values = (
        getattr(shape, "crop_left", 0.0),
        getattr(shape, "crop_right", 0.0),
        getattr(shape, "crop_top", 0.0),
        getattr(shape, "crop_bottom", 0.0),
    )
    return all(abs(float(v)) < 1e-5 for v in values)


def _native_ratio(shape) -> float | None:
    try:
        width_px, height_px = shape.image.size
        if width_px and height_px:
            return float(width_px) / float(height_px)
    except Exception:  # noqa: BLE001 - unsupported picture representation
        return None
    return None


def run_checks(pptx_path: Path) -> dict:
    prs = Presentation(str(pptx_path))
    slide_w = float(prs.slide_width) / EMU_PER_INCH
    slide_h = float(prs.slide_height) / EMU_PER_INCH
    slide_area = slide_w * slide_h

    fails: list[str] = []
    warns: list[str] = []
    slide_summary: list[dict] = []

    for slide_no, slide in enumerate(prs.slides, start=1):
        pictures = [shape for shape in slide.shapes if _is_picture(shape)]
        text_shapes = [shape for shape in slide.shapes if _text(shape)]
        significant = []

        for pic in pictures:
            prect = _rect(pic)
            pname = _shape_name(pic)
            parea = _area(prect)
            ratio_to_slide = parea / slide_area if slide_area else 0.0

            if ratio_to_slide >= SIGNIFICANT_IMAGE_AREA_RATIO:
                significant.append(pic)

            x1, y1, x2, y2 = prect
            if x1 < -1e-4 or y1 < -1e-4 or x2 > slide_w + 1e-4 or y2 > slide_h + 1e-4:
                fails.append(f"p{slide_no}: picture {pname!r} is outside slide bounds")

            native_ratio = _native_ratio(pic)
            if native_ratio and _crop_is_zero(pic):
                shape_ratio = float(pic.width) / float(pic.height) if pic.height else native_ratio
                distortion = abs(shape_ratio / native_ratio - 1.0)
                if distortion > DISTORTION_FAIL_RATIO:
                    fails.append(
                        f"p{slide_no}: picture {pname!r} appears stretched by {distortion:.0%}; preserve aspect ratio or crop intentionally"
                    )
                elif distortion > 0.05:
                    warns.append(
                        f"p{slide_no}: picture {pname!r} aspect ratio differs from source by {distortion:.0%}; verify intentional"
                    )

            if _overlay_allowed(pic):
                continue

            for tshape in text_shapes:
                # Picture objects with captions embedded as text are not text frames in python-pptx,
                # so this only compares native slide text against the picture bbox.
                trect = _rect(tshape)
                overlap = _intersection(prect, trect)
                if overlap > 0.0025:
                    fails.append(
                        f"p{slide_no}: picture {pname!r} overlaps readable text {(_text(tshape)[:45])!r}; reserve separate visual/text zones"
                    )
                    continue

                gap = _edge_gap(prect, trect)
                if gap is None:
                    continue
                if gap < TEXT_GAP_FAIL_IN:
                    fails.append(
                        f"p{slide_no}: picture {pname!r} is only {gap:.2f}in from text {(_text(tshape)[:45])!r}; add safety gap"
                    )
                elif gap < TEXT_GAP_WARN_IN:
                    warns.append(
                        f"p{slide_no}: picture {pname!r} is {gap:.2f}in from text {(_text(tshape)[:45])!r}; recommended gap is >= {TEXT_GAP_WARN_IN:.2f}in"
                    )

        if len(significant) > MAX_SIGNIFICANT_IMAGES:
            warns.append(
                f"p{slide_no}: {len(significant)} significant pictures; classroom slides normally use one hero visual and at most one supporting visual"
            )

        image_area = sum(_area(_rect(pic)) for pic in significant)
        if slide_area and image_area / slide_area > 0.70 and slide_no != 1:
            warns.append(
                f"p{slide_no}: significant pictures occupy about {image_area / slide_area:.0%} of slide area; verify instructional text remains dominant"
            )

        slide_summary.append(
            {
                "slide": slide_no,
                "pictures": len(pictures),
                "significant_pictures": len(significant),
                "text_shapes": len(text_shapes),
            }
        )

    return {
        "file": str(pptx_path),
        "slides": len(prs.slides),
        "fail_count": len(fails),
        "warn_count": len(warns),
        "fails": fails,
        "warnings": warns,
        "slide_summary": slide_summary,
        "note": "Machine QA does not replace final rendered PNG review for contrast, crop quality, or visual dominance.",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Check classroom PPTX visual assets")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", dest="json_path", type=Path)
    args = parser.parse_args()

    if not args.pptx.exists():
        print(f"FAIL: file not found: {args.pptx}")
        return 2
    if args.pptx.suffix.lower() != ".pptx":
        print(f"FAIL: expected .pptx: {args.pptx}")
        return 2

    result = run_checks(args.pptx)
    if args.json_path:
        args.json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    for item in result["fails"]:
        print(f"FAIL: {item}")
    for item in result["warnings"]:
        print(f"WARN: {item}")
    print(f"Visual asset QA: FAIL={result['fail_count']} WARN={result['warn_count']}")
    return 1 if result["fail_count"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
