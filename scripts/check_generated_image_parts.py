#!/usr/bin/env python3
"""Check exact crops and actual PPTX display sizes at the target render size."""
from __future__ import annotations

import argparse
import json
import sys
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def positive_pair(value, label):
    if not isinstance(value, list) or len(value) != 2 or any(type(v) is not int or v <= 0 for v in value):
        raise ValueError(f"{label}: expected two positive integer pixels")
    return value


def check(manifest_path: Path):
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    render_w, render_h = positive_pair(manifest.get("render_px"), "render_px")
    slides = manifest.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("slides: expected a nonempty list")
    base = manifest_path.parent
    prs = Presentation(str(base / manifest["pptx"]))
    if abs(prs.slide_width / prs.slide_height - render_w / render_h) > .01:
        raise ValueError("render_px aspect ratio differs from PPTX canvas")
    numbers = [slide.get("slide_number", idx) for idx, slide in enumerate(slides, start=1)]
    if len(slides) != len(prs.slides) or set(numbers) != set(range(1, len(prs.slides) + 1)):
        raise ValueError("manifest must cover every PPTX slide exactly once")
    errors, checked = [], 0
    for sidx, slide in enumerate(slides, start=1):
        slide_no = slide.get("slide_number", sidx)
        label = f"slide {slide_no}"
        if type(slide_no) is not int or not (1 <= slide_no <= len(prs.slides)):
            errors.append(f"{label}: slide_number not in PPTX")
            continue
        page = prs.slides[slide_no - 1]
        actual = [sh for sh in page.shapes if sh.name.startswith("CLICK")]
        if len(actual) != len({sh.name for sh in actual}):
            errors.append(f"{label}: duplicate CLICK picture names")
        try:
            with Image.open(base / slide["rendered_png"]) as rendered:
                if rendered.size != (render_w, render_h):
                    errors.append(f"{label}: rendered PNG {rendered.size} != render_px {(render_w, render_h)}")
        except (KeyError, OSError) as exc:
            errors.append(f"{label}: missing final rendered PNG: {exc}")
        master_path = base / slide["master"]
        with Image.open(master_path) as master:
            master.load()
            if max(master.size) < 2048:
                errors.append(f"{label}: master {master.size} has long edge below 2048px")
            parts = slide.get("parts")
            if not isinstance(parts, list) or not parts:
                errors.append(f"{label}: parts is empty")
                continue
            names = [part.get("name") for part in parts]
            if len(names) != len(set(names)) or set(names) != {sh.name for sh in actual}:
                errors.append(f"{label}: manifest parts and PPTX CLICK pictures differ")
            for part in parts:
                name = part.get("name", "unnamed")
                pfx = f"{label} {name}"
                rect = part.get("crop_xyxy")
                if (not isinstance(rect, list) or len(rect) != 4
                        or any(type(v) is not int for v in rect)):
                    errors.append(f"{pfx}: invalid crop_xyxy")
                    continue
                x1, y1, x2, y2 = rect
                if not (0 <= x1 < x2 <= master.width and 0 <= y1 < y2 <= master.height):
                    errors.append(f"{pfx}: crop outside master {master.size}")
                    continue
                display = part.get("display_px")
                try:
                    dw, dh = positive_pair(display, f"{pfx} display_px")
                except ValueError as exc:
                    errors.append(str(exc))
                    continue
                with Image.open(base / part["file"]) as tile:
                    tile.load()
                    expected = (x2 - x1, y2 - y1)
                    if tile.size != expected:
                        errors.append(f"{pfx}: tile {tile.size} differs from crop {expected}")
                        continue
                    if tile.convert("RGBA").tobytes() != master.crop(tuple(rect)).convert("RGBA").tobytes():
                        errors.append(f"{pfx}: pixels differ from master crop (resampled or edited)")
                    shapes = [sh for sh in actual if sh.name == name]
                    if not shapes:
                        errors.append(f"{pfx}: picture missing from PPTX")
                        continue
                    shape = shapes[0]
                    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                        errors.append(f"{pfx}: PPTX object is not a picture")
                        continue
                    real_w = shape.width * render_w / prs.slide_width
                    real_h = shape.height * render_h / prs.slide_height
                    if abs(real_w - dw) > 1 or abs(real_h - dh) > 1:
                        errors.append(f"{pfx}: manifest display {dw}x{dh}px differs from PPTX {real_w:.1f}x{real_h:.1f}px")
                    if real_w > tile.width + .01 or real_h > tile.height + .01:
                        errors.append(f"{pfx}: PPTX displays {real_w:.1f}x{real_h:.1f}px, larger than tile {tile.width}x{tile.height}px")
                    if any((shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom)):
                        errors.append(f"{pfx}: PPTX further crops the tile")
                    with Image.open(BytesIO(shape.image.blob)) as embedded:
                        if embedded.convert("RGBA").tobytes() != tile.convert("RGBA").tobytes():
                            errors.append(f"{pfx}: PPTX embedded image differs from crop PNG")
                    if abs(dw * tile.height - dh * tile.width) > max(dw * tile.height, dh * tile.width) * 0.01:
                        errors.append(f"{pfx}: display aspect ratio changes by over 1%")
                    checked += 1
    return checked, errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("manifest", type=Path, help="JSON manifest; image paths are relative to its directory")
    args = parser.parse_args()
    try:
        checked, errors = check(args.manifest)
    except (KeyError, OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"FAIL: {exc}", file=sys.stderr)
        return 1
    for error in errors:
        print(f"FAIL: {error}")
    print(f"IMAGE_PARTS: {'FAIL' if errors else 'PASS'}; checked={checked}; issues={len(errors)}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
