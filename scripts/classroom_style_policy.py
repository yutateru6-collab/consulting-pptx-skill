#!/usr/bin/env python3
"""Shared Classroom Mode style policy helpers.

Classroom text containers are borderless by default. A visible rectangular
outline is allowed only when the boundary itself carries information, such as
a flowchart node, table-like structure, coordinate/axis diagram, Venn diagram,
or an intentional UI mockup.
"""
from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE_TYPE

RECTANGLE_PRESETS = {
    "rect",
    "roundRect",
    "round1Rect",
    "round2SameRect",
    "round2DiagRect",
}

# Semantic boundaries must be named explicitly so the exception is auditable.
SEMANTIC_OUTLINE_PREFIXES = (
    "FLOW_",
    "NODE_",
    "TABLE_",
    "AXIS_",
    "VENN_",
    "DIAGRAM_",
    "UI_",
)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return (shape.text_frame.text or "").strip()


def shape_name(shape) -> str:
    return (getattr(shape, "name", "") or "").strip()


def is_semantic_outline(shape) -> bool:
    name = shape_name(shape).upper()
    return any(name.startswith(prefix) for prefix in SEMANTIC_OUTLINE_PREFIXES)


def _is_rectangular_text_container(shape) -> bool:
    if not shape_text(shape):
        return False

    if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return True

    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
        return False

    try:
        geom = shape._element.spPr.prstGeom
        preset = geom.get("prst") if geom is not None else None
    except Exception:
        return False

    return preset in RECTANGLE_PRESETS


def has_visible_solid_outline(shape) -> bool:
    if not _is_rectangular_text_container(shape):
        return False
    try:
        fill_type = shape.line.fill.type
        return fill_type is not None and str(fill_type).upper().startswith("SOLID")
    except Exception:
        return False


def violates_borderless_policy(shape) -> bool:
    """True when a text container has a decorative rectangular border."""
    return has_visible_solid_outline(shape) and not is_semantic_outline(shape)


def remove_decorative_outline(shape) -> bool:
    """Remove a non-semantic text-container outline. Return True if changed."""
    if not violates_borderless_policy(shape):
        return False
    shape.line.fill.background()
    return True
