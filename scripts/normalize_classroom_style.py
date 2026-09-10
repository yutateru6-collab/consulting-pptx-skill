#!/usr/bin/env python3
"""Normalize a classroom PPTX to the default borderless teaching style.

This removes decorative outlines from text-bearing rectangular containers while
preserving explicitly named semantic boundaries such as FLOW_/NODE_/TABLE_.

Usage:
  python3 scripts/normalize_classroom_style.py deck.pptx --in-place
  python3 scripts/normalize_classroom_style.py deck.pptx -o deck.borderless.pptx
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation

from classroom_style_policy import remove_decorative_outline, shape_name, shape_text


def normalize(input_path: Path, output_path: Path) -> dict:
    prs = Presentation(str(input_path))
    changed = []

    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if remove_decorative_outline(shape):
                changed.append(
                    {
                        "slide": slide_no,
                        "shape": shape_name(shape),
                        "text": shape_text(shape)[:80],
                    }
                )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))

    return {
        "input": str(input_path),
        "output": str(output_path),
        "removed_outline_count": len(changed),
        "changed": changed,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("-o", "--output", type=Path)
    ap.add_argument("--in-place", action="store_true")
    ap.add_argument("--json", type=Path)
    args = ap.parse_args()

    if not args.pptx.exists() or args.pptx.suffix.lower() != ".pptx":
        print(f"FAIL: invalid PPTX path: {args.pptx}")
        return 2

    if args.in_place and args.output:
        print("FAIL: use either --in-place or --output, not both")
        return 2

    if args.in_place:
        output = args.pptx
    elif args.output:
        output = args.output
    else:
        output = args.pptx.with_name(args.pptx.stem + "_borderless.pptx")

    result = normalize(args.pptx, output)
    print(
        f"CLASSROOM_STYLE: removed={result['removed_outline_count']} "
        f"output={result['output']}"
    )

    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
