#!/usr/bin/env python3
"""Smoke tests for the Classroom Mode borderless style policy."""
from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from classroom_style_policy import (
    has_visible_solid_outline,
    is_semantic_outline,
    violates_borderless_policy,
)
from normalize_classroom_style import normalize


def add_rect(slide, name: str, text: str, outlined: bool = True):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(1))
    shape.name = name
    shape.text = text
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    shape.fill.background()
    if outlined:
        shape.line.color.rgb = RGBColor(0, 0, 0)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def main() -> int:
    with TemporaryDirectory() as td:
        td = Path(td)
        src = td / "policy-smoke.pptx"
        out = td / "policy-smoke-normalized.pptx"

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        ordinary = add_rect(slide, "ANIM_01_text", "ordinary text", outlined=True)
        clean = add_rect(slide, "ANIM_02_text", "clean text", outlined=False)
        semantic = add_rect(slide, "FLOW_node_1", "decision node", outlined=True)
        prs.save(src)

        assert has_visible_solid_outline(ordinary)
        assert violates_borderless_policy(ordinary)
        assert not violates_borderless_policy(clean)
        assert is_semantic_outline(semantic)
        assert has_visible_solid_outline(semantic)
        assert not violates_borderless_policy(semantic)

        result = normalize(src, out)
        assert result["removed_outline_count"] == 1, result

        checked = Presentation(out)
        shapes = {s.name: s for s in checked.slides[0].shapes}
        assert not has_visible_solid_outline(shapes["ANIM_01_text"])
        assert not has_visible_solid_outline(shapes["ANIM_02_text"])
        assert has_visible_solid_outline(shapes["FLOW_node_1"])

    print("CLASSROOM_STYLE_SMOKE: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
