#!/usr/bin/env python3
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches, Pt

from check_classroom_textflow import run_checks


def add_text(slide, text, x, y, w, h, size=32):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    return shape


def make_deck(path: Path, bad: bool):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if bad:
        # Narrow + shallow: the English is predicted to wrap to 3 lines and
        # overflow into the Japanese block below.
        add_text(slide, "My brother washes his car every Sunday.", 1.0, 1.5, 3.2, 0.72, 32)
        add_text(slide, "毎週日曜に洗う", 1.0, 2.35, 3.2, 0.55, 26)
    else:
        # Wide + tall + safe gap: same content, but the layout has real reserve.
        add_text(slide, "My brother washes his car every Sunday.", 1.0, 1.5, 6.0, 1.40, 32)
        add_text(slide, "毎週日曜に洗う", 1.0, 3.20, 6.0, 0.70, 26)

    prs.save(path)


def main():
    with TemporaryDirectory() as td:
        td = Path(td)
        bad = td / "bad.pptx"
        good = td / "good.pptx"
        make_deck(bad, True)
        make_deck(good, False)

        bad_result = run_checks(bad)
        good_result = run_checks(good)

        assert bad_result["fail_count"] > 0, bad_result
        assert any("overflow" in x.lower() or "collide" in x.lower() for x in bad_result["fails"]), bad_result
        assert good_result["fail_count"] == 0, good_result

    print("classroom textflow regression tests: PASS")


if __name__ == "__main__":
    main()
