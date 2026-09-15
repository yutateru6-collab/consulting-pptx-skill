#!/usr/bin/env python3
from pathlib import Path
from tempfile import TemporaryDirectory

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from check_classroom_visual_assets import run_checks


def add_text(slide, text, x, y, w, h, size=28):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.clear()
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    return shape


def make_png(path: Path):
    Image.new("RGB", (800, 600), "white").save(path)


def make_deck(path: Path, image_path: Path, bad: bool):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_text(slide, "I have just finished reading this book.", 0.7, 1.5, 6.0, 1.2, 30)
    add_text(slide, "今につながる結果", 0.7, 3.0, 5.0, 0.6, 26)

    if bad:
        # Intentionally invade the text zone.
        slide.shapes.add_picture(str(image_path), Inches(4.8), Inches(1.4), width=Inches(4.0), height=Inches(3.0))
    else:
        # Safe sidecar with clear separation.
        slide.shapes.add_picture(str(image_path), Inches(8.1), Inches(1.4), width=Inches(4.0), height=Inches(3.0))

    prs.save(path)


def main():
    with TemporaryDirectory() as td:
        td = Path(td)
        png = td / "asset.png"
        make_png(png)

        good = td / "good.pptx"
        bad = td / "bad.pptx"
        make_deck(good, png, False)
        make_deck(bad, png, True)

        good_result = run_checks(good)
        bad_result = run_checks(bad)

        assert good_result["fail_count"] == 0, good_result
        assert bad_result["fail_count"] > 0, bad_result
        assert any("overlaps readable text" in x for x in bad_result["fails"]), bad_result

    print("classroom visual asset regression tests: PASS")


if __name__ == "__main__":
    main()
