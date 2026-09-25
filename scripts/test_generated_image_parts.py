#!/usr/bin/env python3
"""Smoke test: exact crops pass; actual PPTX enlargement and altered pixels fail."""
from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from check_generated_image_parts import check


def main():
    with TemporaryDirectory() as td:
        root = Path(td)
        Image.new("RGB", (2048, 1152), (245, 244, 231)).save(root / "master.png")
        with Image.open(root / "master.png") as master:
            master.crop((100, 100, 600, 400)).save(root / "tile.png")
        Image.new("RGB", (1920, 1080), (255, 255, 255)).save(root / "slide.png")

        prs = Presentation()
        prs.slide_width = Inches(13.333333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        sh = slide.shapes.add_picture(
            str(root / "tile.png"), Inches(1), Inches(1),
            width=int(prs.slide_width * 500 / 1920),
            height=int(prs.slide_height * 300 / 1080),
        )
        sh.name = "CLICK01_QUESTION"
        prs.save(root / "deck.pptx")

        manifest = {
            "render_px": [1920, 1080], "pptx": "deck.pptx",
            "slides": [{"slide_number": 1, "master": "master.png", "rendered_png": "slide.png",
                        "parts": [{"name": "CLICK01_QUESTION", "file": "tile.png",
                                   "crop_xyxy": [100, 100, 600, 400],
                                   "display_px": [500, 300], "item_ids": ["q"]}]}]
        }
        path = root / "manifest.json"
        path.write_text(json.dumps(manifest), encoding="utf-8")
        _, issues = check(path)
        assert not issues, issues

        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(root / "deck.pptx")
        try:
            check(path)
        except ValueError as exc:
            assert "every PPTX slide" in str(exc), exc
        else:
            raise AssertionError("a missing slide must not pass the manifest")
        # Restore a one-page file for the enlargement and pixel-mutation cases.
        prs = Presentation()
        prs.slide_width = Inches(13.333333)
        prs.slide_height = Inches(7.5)
        sh = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(
            str(root / "tile.png"), Inches(1), Inches(1),
            width=int(prs.slide_width * 500 / 1920),
            height=int(prs.slide_height * 300 / 1080),
        )
        sh.name = "CLICK01_QUESTION"

        sh.width = int(prs.slide_width * 800 / 1920)
        prs.save(root / "deck.pptx")
        _, issues = check(path)
        assert any("larger than tile" in issue for issue in issues), issues

        sh.width = int(prs.slide_width * 500 / 1920)
        prs.save(root / "deck.pptx")
        Image.new("RGB", (500, 300), (250, 244, 231)).save(root / "tile.png")
        _, issues = check(path)
        assert any("pixels differ from master crop" in issue for issue in issues), issues

    print("GENERATED_IMAGE_PARTS: PASS")


if __name__ == "__main__":
    main()
